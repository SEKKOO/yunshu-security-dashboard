"""Synthetic offline XLSX checks; only the public openpyxl package is used."""
from pathlib import Path
import tempfile
import unittest
import openpyxl
from openpyxl.styles import Font
from render_workbook import render_software, cache_ratio, render_dashboard
from software_inventory import build_model, FAMILIES
from workbook_layout import inspect_template


class PortableTests(unittest.TestCase):
    def test_software_preserves_original_and_lists_versions(self):
        with tempfile.TemporaryDirectory() as d:
            src,out=Path(d)/'source.xlsx',Path(d)/'output.xlsx'
            w=openpyxl.Workbook();s=w.active;s['B1']='原风险区域';s['C2']=10;s['E2']=2;s['G2']='=E2/C2'
            s['B1'].font=Font(bold=True);s.column_dimensions['B'].width=20;s.row_dimensions[1].height=30
            for r,f in zip([5,10,15],FAMILIES):
                s.cell(r,2,f+' 盗版软件安装清单');s.merge_cells(start_row=r,start_column=2,end_row=r,end_column=5)
                for c,v in enumerate(['姓名','部门','设备','安装路径'],2):s.cell(r+1,c,v)
            w.save(src);cache_ratio(src,'G2',0.2)
            raw=[]
            for version in ['2024','2025']:
                for emp in ['e1','e2']:
                    raw.append({'software_name':'Solidworks '+version,'software_version':version,
                                'biz_related_emp':{'id':emp,'name':'同名员工','department':[{'name':'研发'}]},
                                'biz_related_terminal':{'id':emp,'name':emp},'install_path':''})
            snap={'schema':'yunshu-software-v1','queried_at':'2020-01-01T12:00:00+08:00','families':list(FAMILIES),'batches':[{'rows':raw,'total':4}]}
            model=build_model(snap);render_software(src,model,out)
            a=openpyxl.load_workbook(src).active;b=openpyxl.load_workbook(out).active
            for row in a.iter_rows(max_row=4):
                for c in row:self.assertEqual(c.value,b[c.coordinate].value);self.assertEqual(c._style,b[c.coordinate]._style)
            self.assertEqual(openpyxl.load_workbook(out,data_only=True).active['G2'].value,0.2)
            values=[c.value for row in b for c in row if c.value]
            self.assertIn('CATIA 盗版软件安装清单',values);self.assertIn('ABAQUS 盗版软件安装清单',values)
            self.assertEqual(values.count('接口未提供安装路径'),4)
            self.assertEqual(values.count('同名员工'),4)
            self.assertTrue(any('Solidworks 2024' in str(v) for v in values));self.assertTrue(any('Solidworks 2025' in str(v) for v in values))
            # Filled software areas must not be silently cleared on the next invocation.
            with self.assertRaisesRegex(ValueError,'contains data'):render_software(out,model,Path(d)/'bad.xlsx')

    def test_empty_denominator_cache(self):
        template=Path(__file__).resolve().parent.parent/'assets/dashboard-template.xlsx'
        with tempfile.TemporaryDirectory() as d:
            model={'day':[],'month':[],'auditDay':[],'auditMonth':[],'employees':0,'riskEmployees':0,'devices':0,'updated':'2020-01-01','scope':'测试'}
            out=render_dashboard(template,model,inspect_template(template),Path(d))
            self.assertEqual(openpyxl.load_workbook(out,data_only=True).active['G2'].value,'不可计算')
            self.assertTrue(openpyxl.load_workbook(out).active['G2'].value.startswith('=IF('))


if __name__=='__main__':unittest.main()
